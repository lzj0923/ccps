from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"D:\xm\CCPS-Vue3-full-source-20260714")
OUT = ROOT / "docs" / "CCPS系統使用手冊.docx"

BLUE = "2E74B5"
DARK_BLUE = "17365D"
NAVY = "0B315F"
GOLD = "D99400"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "F5F6F8"
MID_GRAY = "D5DCE5"
TEXT_GRAY = "5A6675"
GREEN = "118847"
RED = "C0392B"
PAGE_DXA = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index < len(widths):
                set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_run_font(run, name="Calibri", east_asia="Microsoft JhengHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "更新欄位後顯示"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run)
    add_field(paragraph, "PAGE")
    run = paragraph.add_run(" 頁")
    set_run_font(run)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 18),
        ("Subtitle", 13, TEXT_GRAY, 0, 10),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Figure Caption" not in styles:
        cap = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Figure Caption"]
    cap.font.name = "Calibri"
    cap.font.size = Pt(9)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(9)

    if "Step Number" not in styles:
        step = styles.add_style("Step Number", WD_STYLE_TYPE.PARAGRAPH)
    else:
        step = styles["Step Number"]
    step.base_style = styles["Normal"]
    step.font.name = "Calibri"
    step.font.size = Pt(11)
    step._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    step.paragraph_format.left_indent = Inches(0.28)
    step.paragraph_format.first_line_indent = Inches(-0.24)
    step.paragraph_format.space_after = Pt(5)


def setup_sections(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "CCPS 系統使用手冊  ·  管理端與業主端"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.runs[0]
    set_run_font(hr)
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor.from_string(TEXT_GRAY)

    footer = section.footer
    add_page_number(footer.paragraphs[0])
    for run in footer.paragraphs[0].runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(TEXT_GRAY)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(88)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CCPS")
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(GOLD)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("系統使用手冊")

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("房產、房款、租賃、財務與維修管理操作指南")

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = line.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("適用角色：系統管理員／財務人員／業主")
    set_run_font(r)
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("文件版本：V1.0  ·  2026 年 7 月")
    set_run_font(r)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEXT_GRAY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(95)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("本手冊中的截圖框為預留位置，可於系統正式上線後替換為實際畫面。")
    set_run_font(r)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    doc.add_page_break()


def add_info_table(doc):
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [2300, 7060])
    rows = [
        ("文件名稱", "CCPS 系統使用手冊"),
        ("使用對象", "管理員、財務人員、業主"),
        ("主要範圍", "業主與房產、建案與房款、租賃、財務確認、收支維修、預備金、提醒、同步與報表"),
        ("截圖規則", "每個【截圖位置】框替換為對應畫面；保留圖號與說明文字。"),
    ]
    for index, (key, value) in enumerate(rows):
        table.cell(index, 0).text = key
        table.cell(index, 1).text = value
        set_cell_shading(table.cell(index, 0), LIGHT_BLUE)
        table.cell(index, 0).paragraphs[0].runs[0].bold = True


def add_callout(doc, title, text, tone="blue"):
    fill = PALE_BLUE if tone == "blue" else "FFF6E5" if tone == "gold" else "FDEDEC"
    accent = BLUE if tone == "blue" else GOLD if tone == "gold" else RED
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [PAGE_DXA])
    prevent_row_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 110, 160, 110, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    p = cell.add_paragraph(text)
    p.paragraph_format.space_after = Pt(0)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.add_run(item)


def add_steps(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph(style="Step Number")
        p.add_run(f"{index}. ").bold = True
        p.add_run(item)


def add_screenshot(doc, fig_no, title, scope, height=2.05):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [PAGE_DXA])
    table.style = "Table Grid"
    row = table.rows[0]
    prevent_row_split(row)
    image_path = ROOT / "docs" / "screenshots" / f"{fig_no}.png"
    if image_path.exists():
        cell = table.cell(0, 0)
        set_cell_shading(cell, "FFFFFF")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(image_path), width=Inches(6.25))
        p.paragraph_format.keep_with_next = True
        cap = doc.add_paragraph(style="Figure Caption")
        cap.add_run(f"圖 {fig_no}　{title}（系統實際畫面）")
        return
    row.height = Inches(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"【截圖位置】\n圖 {fig_no}　{title}")
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"建議截圖範圍：{scope}")
    set_run_font(r)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    p.paragraph_format.keep_with_next = True
    cap = doc.add_paragraph(style="Figure Caption")
    cap.add_run(f"圖 {fig_no}　{title}（請以正式環境畫面替換灰色框）")


def add_function(doc, heading, purpose, steps, fig_no, fig_title, fig_scope, notes=None):
    doc.add_heading(heading, level=2)
    p = doc.add_paragraph()
    r = p.add_run("功能說明：")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p.add_run(purpose)
    doc.add_heading("操作步驟", level=3)
    add_steps(doc, steps)
    if notes:
        add_callout(doc, notes[0], notes[1], notes[2] if len(notes) > 2 else "blue")
    add_screenshot(doc, fig_no, fig_title, fig_scope)


def add_chapter(doc, title, intro=None):
    doc.add_page_break()
    doc.add_heading(title, level=1)
    if intro:
        doc.add_paragraph(intro)


def add_role_matrix(doc):
    doc.add_heading("角色與主要職責", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    set_table_width(table, [1700, 2860, 2860, 1940])
    headers = ["角色", "主要操作", "可查看資料", "注意事項"]
    for i, text in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = text
        set_cell_shading(cell, LIGHT_BLUE)
        cell.paragraphs[0].runs[0].bold = True
    set_repeat_table_header(table.rows[0])
    rows = [
        ("系統管理員", "建立業主、建案、房產、租客、租約、工單與預備金資料", "全系統業務資料", "新增與修改前先核對關聯對象"),
        ("財務人員", "確認房款、租金、預備金充值；查閱憑證與同步狀態", "待確認與歷史交易", "確認後會影響入帳及餘額"),
        ("業主", "查看名下房產、房款進度、租金收入、收支、預備金、通知與文件", "本人及名下單位資料", "付款憑證依畫面規則上傳"),
    ]
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = value


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings._element.append(update_fields)
    setup_styles(doc)
    setup_sections(doc)
    add_cover(doc)

    doc.add_heading("手冊說明", level=1)
    doc.add_paragraph("本手冊以業務流程為主線，說明 CCPS 系統中管理端與業主端的實際操作方式。建議新使用者先閱讀「快速上手」，再依工作角色查閱相關章節。")
    add_info_table(doc)
    add_callout(doc, "重要原則", "凡涉及金額確認的事項，應由管理端在「財務確認」中處理；不需要確認的資料則保留於歷史記錄。租金為線下完成收款後，由管理員在線上登記、確認及上傳或更換憑證。", "gold")

    doc.add_heading("目錄", level=1)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_paragraph("提示：在 Microsoft Word 中按 Ctrl+A，再按 F9，可更新目錄、頁碼及圖表欄位。", style="Figure Caption")

    add_role_matrix(doc)

    add_chapter(doc, "1　快速上手與共通操作", "本章先說明登入、導覽、搜尋篩選及一般資料操作。不同模組的畫面結構一致，可用相同方式快速上手。")
    add_function(doc, "1.1　登入與切換入口", "依帳號角色進入管理端或業主端；系統只顯示該角色可使用的功能。",
                 ["開啟 CCPS 登入頁。", "輸入帳號及密碼。", "按「登入」進入對應工作台。", "確認右上角顯示的使用者名稱與角色正確。"],
                 "1-1", "系統登入頁", "登入表單、角色提示與登入按鈕")
    add_function(doc, "1.2　主選單與頁面導覽", "左側主選單用於切換業務模組；上方工具列提供全域搜尋、日期區間、通知與使用者入口。",
                 ["從左側選單選擇目標模組。", "利用上方搜尋框輸入業主、單位、租客或文件關鍵字。", "需要按期間查詢時設定開始與結束日期。", "完成操作後可直接切換其他模組，資料會重新載入。"],
                 "1-2", "管理端主介面", "左側選單、頁首搜尋、日期及使用者區域")
    add_function(doc, "1.3　搜尋、篩選與分頁", "清單資料由資料庫分頁查詢，搜尋與篩選會共同作用。房產頁的狀態篩選以出租狀態為主。",
                 ["輸入搜尋關鍵字。", "選擇建案或出租狀態。", "按頁碼或上一頁／下一頁查看其他資料。", "清除關鍵字並選回「全部」即可恢復完整清單。"],
                 "1-3", "清單搜尋與分頁", "搜尋框、建案篩選、出租狀態篩選與頁碼")

    add_chapter(doc, "2　買房業務完整流程", "買房資料必須依順序建立關聯：建案 → 業主 → 房產／單位 → 購房合約 → 房款計劃 → 付款憑證 → 財務確認。")
    add_callout(doc, "流程總覽", "新增建案 → 新增業主 → 新增房產並綁定業主 → 建立購房合約與房款分期 → 業主提交付款憑證 → 財務確認 → 更新付款進度與交房狀態", "blue")
    add_function(doc, "2.1　新增建案", "先建立建案主檔，供房產、房款、租約、維修及報表引用。",
                 ["進入「建案與房款」。", "按「新增建案」。", "填寫建案編碼、建案名稱、地址、城市及國家代碼。", "選擇啟用狀態並按「確認新增」。"],
                 "2-1", "新增建案", "建案與房款頁及新增建案視窗",
                 ("欄位要求", "建案編碼與名稱為必填；建案編碼應使用公司統一規則，建立後避免隨意更改。", "gold"))
    add_function(doc, "2.2　新增業主", "建立購房人資料，後續房產、購房合約、房款及預備金均以業主為關聯對象。",
                 ["進入「業主管理」。", "按「新增業主」。", "填寫繁體中文姓名、英文姓名、證件、電話、電郵及通訊資料。", "確認資料無誤後儲存。"],
                 "2-2", "新增業主資料", "業主管理清單與新增業主視窗")
    add_function(doc, "2.3　新增房產並綁定業主", "在房產管理中建立實際房屋／單位，並設定建案、單位號、房產階段及業主關聯。",
                 ["進入「房產管理」。", "按「新增房產」。", "選擇建案並填寫單位號、房型、面積、售價等資料。", "選擇業主並建立購房關聯。", "依實際情況設定「未交房」或「已交房待出租」等階段後儲存。"],
                 "2-3", "新增房產與綁定業主", "房產管理、新增房產表單及業主選擇欄位",
                 ("房產狀態", "「未交房」代表仍處於購房／建設階段；交房後應改為「已交房待出租」，出租後再更新為「出租中」。", "gold"))
    add_function(doc, "2.4　建立房款計劃與分期", "依購房合約總價建立工程進度付款計劃，設定每一期的工程階段、到期日及應收金額。",
                 ["回到「建案與房款」並按「新增房款」。", "選擇尚未建立有效付款計劃的購房合約。", "輸入計劃名稱與生效日期。", "按「增加一期」建立各期資料，或使用「平均分期」。", "確認各期到期日依序遞增，且分期合計等於合約總價。", "按「確認建立房款」。"],
                 "2-4", "新增房款計劃", "購房合約選擇、分期編輯器及合計校驗區域")
    add_function(doc, "2.5　維護本期房款與催繳", "在房款進度清單中查看各期應收、已收、未收、到期日及狀態，必要時修正工程階段／到期日或發送提醒。",
                 ["在清單中選擇目標建案及單位。", "選中需處理的分期。", "使用「編輯本期房款」修改工程階段或到期日。", "逾期或即將到期時使用催繳／提醒功能。", "已有收款時可查看本期付款記錄。"],
                 "2-5", "房款進度與本期操作", "房款清單、右側詳情及本期操作按鈕")
    add_function(doc, "2.6　業主提交房款憑證", "業主端按分期上傳付款憑證，供財務核對銀行入帳資料。",
                 ["業主登入後進入「房款進度」。", "選擇名下房產及待付款期數。", "按「上傳付款憑證」。", "選擇付款方式、付款日期、金額及憑證文件後提交。", "回到清單確認狀態已變為待財務確認。"],
                 "2-6", "業主上傳付款憑證", "房款進度、期數詳情與憑證上傳區域")
    add_function(doc, "2.7　財務確認房款", "財務人員僅在待確認清單中處理需要核對的房款；完成或退回的交易進入歷史記錄。",
                 ["管理端進入「財務確認」，選擇「買房款」。", "在待確認清單選擇交易並查看付款憑證。", "核對付款人、銀行參考、付款日期、金額及對應期數。", "資料一致時按「確認收款」，填寫審核備註並入帳。", "資料不完整時按「退回補件」，填寫明確退回原因。", "需要查看已處理資料時按「歷史記錄」，完成後按「返回財務確認」。"],
                 "2-7", "買房款財務確認", "買房款頁籤、待確認交易、右側詳情與確認／退回按鈕",
                 ("入帳影響", "確認收款後，金額會累加至對應房款期數並進入待同步會計狀態；退回不等於完成，退回原因會保留。", "gold"))

    add_chapter(doc, "3　房產與出租管理", "房主與房產分開管理：業主管理展示所有房主；房產管理展示所有房屋。房產清單可按出租狀態篩選。")
    add_function(doc, "3.1　查看與編輯業主", "集中查看全部業主、聯絡資料及其名下房產，並維護資料。",
                 ["進入「業主管理」。", "搜尋業主姓名、電話或證件資料。", "選擇一筆業主查看右側詳情。", "按編輯功能更新資料，儲存後重新確認。"],
                 "3-1", "業主管理", "業主清單、搜尋與右側業主詳情")
    add_function(doc, "3.2　查看與編輯房產", "集中管理全部房子與單位，並以資料庫分頁方式查詢。",
                 ["進入「房產管理」。", "以建案、單位號或業主搜尋。", "以出租狀態篩選「未交房」、「已交房待出租」、「出租中」或「未啟用出租」。", "選擇房產查看詳情，必要時編輯房產階段、出租狀態或業主關聯。"],
                 "3-2", "房產管理與出租狀態", "房產清單、出租狀態篩選、分頁及詳情")

    add_chapter(doc, "4　租客、租約與收租流程", "本系統不需要另建租戶端審批流程。線下收款完成後，由管理員在系統內確認租金並保存憑證；未付或部分支付的租約都可進行收租確認。")
    add_callout(doc, "流程總覽", "已交房待出租 → 新增租客 → 新增／上傳租約 → 生成租金帳單 → 管理員確認收款 → 上傳或更換憑證 → 業主端查看租金收入", "blue")
    add_function(doc, "4.1　新增租客", "建立承租人基本資料，作為租約及租金帳單的付款人。",
                 ["進入「租客與租金」。", "按「新增租客」。", "填寫繁體姓名、證件、電話、電郵及聯絡地址。", "儲存後在租客清單確認資料。"],
                 "4-1", "新增租客", "租客與租金頁及新增租客視窗")
    add_function(doc, "4.2　新增租約並上傳文件", "將租客與可出租單位建立租賃關係，設定租期、月租、押金及付款日，並在同一頁保存租約文件。",
                 ["按「新增租約」。", "選擇租客與「已交房待出租」的單位。", "填寫租約編號、開始／結束日期、月租、押金及每月到期日。", "上傳 PDF、JPG 或 PNG 租約文件。", "儲存後按「查看租約」確認文件可正常開啟。"],
                 "4-2", "新增與查看租約", "新增租約表單、文件上傳及查看租約按鈕",
                 ("資料關聯", "租約啟用後，單位出租狀態應同步為「出租中」，並依租約產生租金帳單。", "gold"))
    add_function(doc, "4.3　編輯租約", "租期、租金、到期日、文件或租客資料有變動時，可更新現有租約。",
                 ["在租客與租金清單選擇目標租約。", "開啟操作選單並選擇「編輯租約」。", "修改允許調整的租期、金額或文件。", "儲存並確認後續帳單資料正確。"],
                 "4-3", "編輯租約", "租約操作選單與編輯視窗")
    add_function(doc, "4.4　轉租", "原租約需轉由新租客承租時，以轉租功能保留原租約歷史並建立新的承租關係。",
                 ["選擇原租約並開啟「轉租」。", "選擇或新增接手租客。", "輸入轉租生效日及新租約條件。", "核對未結清租金、押金及原租約終止日。", "確認轉租後查看新舊租約狀態。"],
                 "4-4", "租約轉租", "轉租視窗、新租客及新舊租約關聯資料",
                 ("操作注意", "轉租前應先處理原租約未結清款項；不要直接覆蓋原租客資料，以免歷史帳單失去追溯性。", "gold"))
    add_function(doc, "4.5　確認未付或部分支付租金", "租金頁應展示所有未支付及未完全支付的租約帳單，管理員可直接登記本次實收。",
                 ["進入「租客與租金」，篩選待收、逾期或部分收款。", "選擇租約並按「確認租金／收租」。", "輸入本次實收金額、付款日期、付款方式及銀行參考。", "需要留存時上傳付款憑證。", "確認後檢查本期已收、未收及狀態是否更新。", "部分收款可再次操作，直至本期全額收清。"],
                 "4-5", "租金收款確認", "未付／部分支付帳單、收款視窗與金額欄位",
                 ("業務規則", "租金不走租戶提交及線上審批流程；管理員依線下已確認的收款結果直接登記。憑證可於後續更換。", "blue"))
    add_function(doc, "4.6　查看租金歷史與更換憑證", "已完成的租金交易保留於歷史記錄，可查看租金帳單、付款資料、SQL 同步狀態及補上／更換憑證。",
                 ["在財務確認選擇「租金」。", "按「歷史記錄」查看已處理租金。", "選擇交易查看關聯租約與帳單月份。", "按憑證操作補上或更換 PDF、JPG、PNG 文件。", "按返回按鈕回到待確認頁。"],
                 "4-6", "租金歷史與憑證", "租金歷史清單、右側詳情及更換憑證入口")

    add_chapter(doc, "5　財務確認管理", "財務確認頁只顯示需要人員確認的項目，分為買房款、租金與預備金充值三類；已確認或已退回資料統一放在歷史記錄。")
    add_function(doc, "5.1　待確認與歷史記錄", "利用單一「歷史記錄」按鈕在待辦與已處理資料之間切換，保持頁面簡潔。",
                 ["進入「財務確認」。", "選擇買房款、租金或預備金充值。", "預設查看待確認項目。", "按「歷史記錄」查看已確認／已退回資料。", "按「返回財務確認」回到待辦清單。"],
                 "5-1", "財務確認與歷史切換", "三類頁籤、歷史記錄按鈕與返回按鈕")
    add_function(doc, "5.2　批量確認", "當多筆交易均已完成銀行核對，可勾選後批量確認，減少重複操作。",
                 ["在待確認清單勾選多筆交易。", "按「批量確認」。", "輸入批量審核備註。", "再次確認筆數與金額後提交。", "刷新清單並抽查歷史記錄。"],
                 "5-2", "批量財務確認", "清單勾選框、批量確認按鈕與備註視窗",
                 ("風險提示", "只可批量確認已逐筆核對的交易。預備金批量確認採同一事務，任一筆校驗失敗時整批不入帳。", "red"))

    add_chapter(doc, "6　預備金管理", "預備金用於記錄業主／單位的最低標準、當前餘額、充值、扣款及不足提醒。管理端可設定標準並直接充值，也可確認待入帳的充值申請。")
    add_function(doc, "6.1　設定最低預備金", "為每個業主／單位設定最低應保留餘額，系統據此判斷正常、餘額不足或待充值。",
                 ["進入「預備金」。", "選擇業主／單位。", "按「設定預備金」。", "輸入最低標準及備註。", "儲存後檢查當前狀態是否正確。"],
                 "6-1", "設定最低預備金", "預備金清單、右側帳戶資料與設定視窗")
    add_function(doc, "6.2　管理端直接充值", "管理員已在線下確認款項時，可直接為指定預備金帳戶充值並建立交易紀錄。",
                 ["在預備金清單選擇帳戶。", "按「直接充值／新增充值」。", "輸入金額、付款方式、付款日期、銀行參考及備註。", "需要時上傳憑證。", "確認後檢查餘額與累計充值是否增加。"],
                 "6-2", "預備金直接充值", "帳戶詳情、充值按鈕及充值視窗",
                 ("使用時機", "直接充值適用於管理員已完成線下核對的情況；需要財務二次核對的充值應進入財務確認。", "blue"))
    add_function(doc, "6.3　確認預備金充值", "財務確認後才將待審充值加入帳戶餘額；退回時必須保留原因。",
                 ["進入「財務確認」並選擇「預備金充值」。", "選擇待確認交易並查看付款資料及憑證。", "資料正確時按「確認充值」，填寫備註後入帳。", "資料錯誤時按「退回申請」，填寫退回原因。", "到歷史記錄檢查審核人、時間、備註及同步狀態。"],
                 "6-3", "預備金充值財務確認", "待確認充值、帳戶餘額、確認後餘額及審核按鈕")
    add_function(doc, "6.4　查看充值、扣款與餘額不足", "追蹤每次預備金異動，並對低於最低標準的帳戶進行補繳提醒。",
                 ["選擇帳戶查看交易明細。", "核對充值、扣款、關聯工單及操作人。", "篩選餘額不足或待充值帳戶。", "依公司流程通知業主補足預備金。"],
                 "6-4", "預備金流水與餘額狀態", "交易明細、最低標準、目前餘額及不足狀態")

    add_chapter(doc, "7　收支與維修管理", "收支與維修資料直接連接資料庫。維修工單建立後，可從預備金扣款；若餘額不足則由管理員處理款項，並上傳維修前後圖片後完成工單。")
    add_function(doc, "7.1　新增支出", "登記管理費、水電費、維修費或其他支出，保存單位、供應商、金額、支付方式與憑證。",
                 ["進入「收支與維修」。", "按「新增支出」。", "選擇建案／單位及支出類別。", "填寫說明、供應商、日期、金額與支付方式。", "選擇是否從預備金扣除並上傳附件。", "儲存後核對清單及業主端顯示。"],
                 "7-1", "新增支出", "新增支出視窗、預備金扣除選項與附件欄位")
    add_function(doc, "7.2　新增維修工單", "建立維修事項，記錄報修時間、類別、內容、預估費用及初始附件。",
                 ["按「新增維修」。", "選擇建案／單位及維修類別。", "填寫事項、報修時間、費用、供應商與處理說明。", "上傳報修或維修前照片。", "儲存後工單狀態應為待處理或處理中。"],
                 "7-2", "新增維修工單", "新增維修表單、類別、費用及附件區域")
    add_function(doc, "7.3　處理維修與預備金扣款", "使用清單中的「處理」按鈕進入工單，選擇從預備金扣除或記錄餘額不足的後續處理。",
                 ["在維修工單清單找到處理中項目。", "按「處理」。", "確認實際費用及支付來源。", "預備金充足時確認扣款；不足時依線下收款結果補充處理資料。", "保存處理進度。"],
                 "7-3", "維修工單處理", "維修工單列、處理按鈕及預備金扣款區域",
                 ("餘額檢查", "扣款前必須核對當前預備金餘額與最低標準，並確保工單、支出及預備金流水使用同一筆關聯資料。", "gold"))
    add_function(doc, "7.4　上傳維修前後圖片並完成", "工單完成前需保留維修前及維修後圖片，確認實際費用後再將狀態改為已完成。",
                 ["打開工單處理視窗。", "分別上傳「維修前」及「維修後」圖片。", "核對圖片、附件數量、實際費用及處理說明。", "按「確認完成」。", "回到清單確認狀態為已完成，且已完成工單不再顯示處理按鈕。"],
                 "7-4", "維修附件與完成工單", "維修前／後圖片區、實際費用與確認完成按鈕")

    add_chapter(doc, "8　自動提醒、會計同步與報表", "本章說明日常輔助功能：提醒到期事項、將交易同步至 SQL Account，以及產生管理報表。")
    add_function(doc, "8.1　自動提醒", "針對房款、租金、預備金、維修及文件到期設定通知規則並查看發送結果。",
                 ["進入「自動提醒」。", "按「新增提醒」建立場景、對象、觸發條件及通知方式。", "啟用規則或使用立即／批量發送。", "查看成功、失敗及失敗原因，必要時重發。"],
                 "8-1", "自動提醒", "提醒規則、通知方式、發送狀態及重發操作")
    add_function(doc, "8.2　SQL Account 同步", "將業主／租客、租金帳單、房款、維修支出、預備金流水及憑證同步至會計系統。",
                 ["進入「SQL Account 同步」。", "查看待同步、同步中、已同步及同步失敗批次。", "按「立即同步」處理待同步資料。", "失敗時查看錯誤原因，修正來源資料後重試。", "需要稽核時下載日誌或比較差異。"],
                 "8-2", "SQL Account 同步", "同步批次、來源資料、成功／失敗筆數及錯誤原因")
    add_function(doc, "8.3　報表與導出", "按期間與建案產生房款、租金、收支維修、預備金、財務確認及同步紀錄。",
                 ["進入「報表與導出」。", "選擇報表類型、期間、建案及狀態。", "按「新增報表」產生資料。", "預覽內容後以 PDF 或 XLSX 匯出。", "批量導出時確認檔案範圍及下載結果。"],
                 "8-3", "報表與導出", "報表類型、篩選條件、預覽與下載操作")

    add_chapter(doc, "9　業主端功能", "業主端以查看本人及名下單位資料為主，協助掌握房款、租金收入、收支、預備金、通知及文件。")
    owner_functions = [
        ("9.1　我的房產", "查看名下全部房產、建案、單位、房產階段及各模組摘要。", ["登入業主端。", "進入「我的房產」。", "選擇單位查看房款、租金、預備金及文件摘要。"], "9-1", "業主端房產總覽", "名下房產清單與右側摘要"),
        ("9.2　房款進度", "查看每期應繳、已繳、未繳、到期日及財務確認結果，並上傳付款憑證。", ["進入「房款進度」。", "選擇房產與期數。", "查看付款進度，必要時上傳憑證。", "被退回時依退回原因補正並重新提交。"], "9-2", "業主房款進度", "付款總覽、分期明細、退回原因及上傳入口"),
        ("9.3　租金收入", "查看租約、每月租金、已收／未收、付款日期及租金歷史。", ["進入「租金收入」。", "選擇出租單位。", "查看本月收入、待收租金與租約到期。", "使用下載功能保存租金明細。"], "9-3", "業主租金收入", "租金摘要、租約資料及歷史清單"),
        ("9.4　收支與維修", "查看管理端登記的收入、支出、維修費、附件及工單狀態。", ["進入「收支與維修」。", "按單位或類別篩選。", "選擇資料查看費用、憑證及維修前後附件。"], "9-4", "業主收支維修", "收支清單、工單狀態及附件"),
        ("9.5　預備金", "查看當前餘額、最低標準、充值、扣款及低餘額提醒。", ["進入「預備金」。", "選擇單位查看帳戶。", "查看交易明細及關聯維修扣款。", "需要補繳時依頁面提示提交或聯絡管理員。"], "9-5", "業主預備金", "餘額、最低標準、充值扣款及提醒"),
        ("9.6　通知中心", "集中查看房款到期、租金、文件、預備金及維修通知。", ["進入「通知中心」。", "篩選未讀或通知類型。", "打開通知查看關聯資料。", "處理後標記已讀。"], "9-6", "業主通知中心", "通知分類、未讀狀態及內容"),
        ("9.7　文件中心", "查看或下載買賣合約、租約、付款憑證及其他文件，追蹤待補件與到期狀態。", ["進入「文件中心」。", "按單位或文件類型篩選。", "查看文件狀態與到期日。", "按查看／下載保存文件；需要時上傳補充資料。"], "9-7", "業主文件中心", "文件分類、狀態、到期日及查看／下載操作"),
    ]
    for args in owner_functions:
        add_function(doc, *args)

    add_chapter(doc, "10　狀態說明與常見問題", "狀態名稱直接影響下一步操作。遇到資料不一致時，先確認目前狀態，再檢查關聯對象、金額及附件。")
    doc.add_heading("10.1　常用狀態", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_width(table, [2100, 2860, 4400])
    for i, text in enumerate(["狀態", "適用範圍", "處理方式"]):
        table.cell(0, i).text = text
        set_cell_shading(table.cell(0, i), LIGHT_BLUE)
        table.cell(0, i).paragraphs[0].runs[0].bold = True
    set_repeat_table_header(table.rows[0])
    statuses = [
        ("待確認", "房款／充值", "財務核對憑證及銀行資料後確認或退回。"),
        ("已確認", "各類款項", "已完成審核；到歷史記錄查看。"),
        ("已退回", "房款／充值", "查看退回原因，修正後重新提交；不可視為已完成。"),
        ("部分收款", "租金", "可再次登記收款，直至本期未收為零。"),
        ("逾期", "房款／租金", "確認到期日及未收金額，發送提醒或進行收款。"),
        ("未交房", "房產", "仍在購房／建設階段，通常不可新增租約。"),
        ("已交房待出租", "房產", "已可建立租約。"),
        ("出租中", "房產", "已有有效租約。"),
        ("待同步／同步失敗", "會計同步", "查看來源資料及失敗原因，修正後重試。"),
    ]
    for row_data in statuses:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = value

    doc.add_heading("10.2　常見問題處理", level=2)
    faqs = [
        ("找不到可建立房款的合約", "確認已建立業主、房產及購房合約，且該合約尚未存在有效付款計劃。"),
        ("新增房款時無法儲存", "檢查每期金額是否大於零、到期日是否依序遞增，以及分期合計是否等於合約總價。"),
        ("租約無法查看", "確認租約文件已在租客與租金頁上傳，格式為系統支援的 PDF／JPG／PNG。"),
        ("已退回卻顯示完成", "退回交易應顯示「已退回」並保留原因；如狀態異常，重新載入後聯絡系統管理員檢查資料。"),
        ("租金沒有租戶提交入口", "這是既定設計：租金在線下確認後由管理員直接登記，不需要租戶端審批。"),
        ("維修無法從預備金扣款", "核對帳戶餘額、最低標準及工單實際費用；不足時先按公司流程處理補款。"),
        ("SQL 同步失敗", "查看失敗原因與來源批次，修正缺失欄位、日期、科目映射或附件路徑後重試。"),
    ]
    for question, answer in faqs:
        p = doc.add_paragraph()
        r = p.add_run(f"問：{question}")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        p = doc.add_paragraph(f"答：{answer}")
        p.paragraph_format.left_indent = Inches(0.22)

    add_chapter(doc, "附錄 A　截圖製作與替換清單", "建議在正式測試環境使用相同瀏覽器寬度截圖，並避免顯示真實證件、電話、銀行帳號或其他個人資料。")
    add_bullets(doc, [
        "截圖前先準備一組完整示範資料，名稱使用繁體中文或正式英文姓名，不使用 ADMIN Test 等測試字樣。",
        "保留頁面標題、主要按鈕、表格欄位及右側詳情；不相關區域可裁切。",
        "若畫面含個人資料，請先遮蔽證件號、電話、電郵、銀行參考及文件內容。",
        "將截圖貼入灰色框後，維持原有圖號與圖說；必要時等比例縮放，避免拉伸變形。",
        "全部替換後在 Word 中按 Ctrl+A、F9 更新目錄，再檢查分頁及圖說。",
    ])
    doc.add_heading("建議優先截圖的核心流程", level=2)
    core = [
        ("買房", "圖 2-1 至圖 2-7", "建案、業主、房產、房款、憑證與財務確認"),
        ("出租", "圖 4-1 至圖 4-6", "租客、租約、轉租、收租與歷史憑證"),
        ("預備金", "圖 6-1 至圖 6-4", "標準、直接充值、財務確認與流水"),
        ("維修", "圖 7-1 至圖 7-4", "支出、工單、扣款、前後圖片與完成"),
        ("業主端", "圖 9-1 至圖 9-7", "房產、房款、租金、收支、預備金、通知與文件"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_width(table, [1800, 2200, 5360])
    for i, value in enumerate(["流程", "圖號", "內容"]):
        table.cell(0, i).text = value
        set_cell_shading(table.cell(0, i), LIGHT_BLUE)
        table.cell(0, i).paragraphs[0].runs[0].bold = True
    for row_data in core:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = value

    doc.add_paragraph()
    add_callout(doc, "版本維護", "系統功能、欄位或流程調整後，應同步更新本手冊的操作步驟、截圖、狀態說明及版本日期。", "blue")

    doc.save(OUT)
    print(OUT)
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} screenshot_placeholders={sum('【截圖位置】' in cell.text for table in doc.tables for row in table.rows for cell in row.cells)}")


if __name__ == "__main__":
    build()
