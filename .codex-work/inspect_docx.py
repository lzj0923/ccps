from pathlib import Path
import sys
from docx import Document


work_dir = Path(__file__).parent
path = Path(sys.argv[1]) if len(sys.argv) > 1 else max(work_dir.glob("*.docx"), key=lambda candidate: candidate.stat().st_size)
doc = Document(path)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} inline_shapes={len(doc.inline_shapes)} sections={len(doc.sections)}")
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"P{index:03d} [{paragraph.style.name}] {text}")

for table_index, table in enumerate(doc.tables):
    print(f"TABLE {table_index} rows={len(table.rows)} cols={len(table.columns)}")
    for row_index, row in enumerate(table.rows):
        cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        print(f"  R{row_index:03d}: " + " || ".join(cells))
