from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORK_DIR = Path(__file__).parent
OUTPUT = WORK_DIR.parent / "CCPS_今日頁面功能交付說明_優化版.docx"
IMAGE_DIR = WORK_DIR / "source_images"
LOGO = WORK_DIR.parent / "frontend" / "ccps-logo.png"

# standard_business_brief preset with a named CCPS brand-color override.
NAVY = "082B57"
TEAL = "008D95"
GOLD = "E6A900"
INK = "1D2A3A"
MUTED = "6B778C"
LIGHT = "F2F4F7"
PALE_TEAL = "EAF7F7"
BORDER = "D8DEE8"
WHITE = "FFFFFF"
GREEN = "1F8A55"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft JhengHei")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        node = tbl_pr.find(qn(tag))
        if node is not None:
            tbl_pr.remove(node)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, (cell, width) in enumerate(zip(row.cells, widths_dxa)):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, size=9.4, bold=False, color=INK):
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    suffix = paragraph.add_run(" 頁")
    set_run_font(suffix, size=9, color=MUTED)


def add_paragraph(doc, text="", size=11, color=INK, bold=False, italic=False,
                  align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.1):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def create_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), "bullet")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•")
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, indent])
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(fonts)
    level.extend([start, number_format, level_text, level_justification, p_pr, r_pr])
    abstract.append(level)
    first_num_index = next((index for index, child in enumerate(numbering) if child.tag == qn("w:num")), len(numbering))
    numbering.insert(first_num_index, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_heading(doc, text, level, page_break_before=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.page_break_before = page_break_before
    run = p.add_run(text)
    size = {1: 16, 2: 13, 3: 12}[level]
    color = TEAL if level in (1, 2) else NAVY
    set_run_font(run, size=size, color=color, bold=True)
    return p


def add_bullet(doc, text, compact=False):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.keep_together = True
    if compact:
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.12
    p_pr = p._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(BULLET_NUM_ID))
    num_pr.extend([ilvl, num_id])
    p_pr.insert(0, num_pr)
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_labeled_note(doc, label, text, fill=PALE_TEAL, compact=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3 if compact else 6)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.0 if compact else 1.1
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, size=9.8 if compact else 10.2, color=TEAL, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=9.8 if compact else 10.2, color=INK)
    return p


def add_screenshot(doc, number, title, alt_text, width=6.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(IMAGE_DIR / f"image-{number}.png"), width=Inches(width))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    caption = add_paragraph(doc, f"圖 {number}　{title}", size=9.3, color=MUTED,
                            italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=5, line=1.0)
    caption.paragraph_format.keep_with_next = True


def start_new_page_section(doc):
    new_section = doc.add_section(WD_SECTION.NEW_PAGE)
    new_section.different_first_page_header_footer = False
    new_section.header.is_linked_to_previous = True
    new_section.footer.is_linked_to_previous = True
    return new_section


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
normal.font.size = Pt(11)
normal.font.color.rgb = rgb(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for style_name, size, color, before, after in (
        ("Heading 1", 16, TEAL, 16, 8),
        ("Heading 2", 13, TEAL, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4)):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = rgb(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True

bullet_style = styles["List Bullet"]
bullet_style.font.name = "Calibri"
bullet_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
bullet_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
bullet_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
bullet_style.font.size = Pt(11)
bullet_style.paragraph_format.left_indent = Inches(0.5)
bullet_style.paragraph_format.first_line_indent = Inches(-0.25)
bullet_style.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))
bullet_style.paragraph_format.space_after = Pt(8)
bullet_style.paragraph_format.line_spacing = 1.167
BULLET_NUM_ID = create_bullet_numbering(doc)

header_p = section.header.paragraphs[0]
header_p.paragraph_format.space_after = Pt(0)
left = header_p.add_run("CCPS 房產管理系統")
set_run_font(left, size=9, color=TEAL, bold=True)
right = header_p.add_run("　｜　頁面功能交付說明")
set_run_font(right, size=9, color=MUTED)
add_page_number(section.footer.paragraphs[0])
section.first_page_header.paragraphs[0].text = ""
add_page_number(section.first_page_footer.paragraphs[0])

# Cover: editorial_cover pattern.
add_paragraph(doc, "CCPS PROPERTY MANAGEMENT", size=10, color=GOLD, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, before=32, after=14)
logo_p = doc.add_paragraph()
logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo_p.paragraph_format.space_after = Pt(24)
logo_shape = logo_p.add_run().add_picture(str(LOGO), width=Inches(2.1))
logo_shape._inline.docPr.set("descr", "CCPS 家慶企業標誌")
add_paragraph(doc, "今日頁面功能交付說明", size=28, color=NAVY, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=9, line=1.0)
add_paragraph(doc, "頁面展示內容、互動功能與後端整合", size=14, color=TEAL,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=28, line=1.0)
add_paragraph(doc, "業主端核心頁面｜版本 1.0｜2026 年 7 月 15 日", size=10.5,
              color=MUTED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=34)
add_labeled_note(doc, "文件目的", "整理今日完成的 5 個核心頁面，說明畫面所呈現的資訊、可操作功能、資料串接與目前實現邊界，作為交付、驗收及後續迭代依據。")

start_new_page_section(doc)
add_heading(doc, "文件範圍與交付摘要", level=1)
add_paragraph(doc, "本次整理涵蓋登入、房產總覽、房款進度、付款憑證上傳及租金收入 5 個頁面。所有說明均依據目前前後端程式與原始頁面截圖整理。")

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
header_flag = OxmlElement("w:tblHeader")
header_flag.set(qn("w:val"), "true")
header_tr_pr.append(header_flag)
headers = ["頁面", "主要目的", "核心輸出", "後端整合"]
for index, text in enumerate(headers):
    set_cell_shading(table.rows[0].cells[index], LIGHT)
    set_cell_text(table.rows[0].cells[index], text, size=9.2, bold=True, color=NAVY)
rows = [
    ("登入頁", "建立安全登入入口", "帳號驗證、工作階段", "登入／登出／Session"),
    ("我的房產", "總覽業主持有單位", "KPI、房產清單、通知", "業主儀表板"),
    ("房款進度", "追蹤購屋付款計畫", "期數、到期、狀態、摘要", "付款進度查詢"),
    ("上傳付款憑證", "提交付款資料供財務審核", "表單、附件、提交記錄", "Multipart 憑證提交"),
    ("租金收入", "掌握應收與到帳情況", "KPI、明細、趨勢、匯出", "租金收入查詢"),
]
for row_values in rows:
    cells = table.add_row().cells
    for index, text in enumerate(row_values):
        set_cell_text(cells[index], text, size=9.0, bold=index == 0, color=TEAL if index == 0 else INK)
set_table_geometry(table, [1700, 2280, 2860, 2520])

add_heading(doc, "整體交付成果", level=2)
for text in (
    "完成業主端頁面導覽、帳號選單、通知入口與頁面狀態切換，形成一致的操作框架。",
    "核心資料頁已串接後端 API，包含登入工作階段、業主儀表板、房款進度、付款憑證提交及租金收入。",
    "主要資料畫面具備載入中、錯誤、空資料或重新載入等狀態，避免只呈現靜態樣板。",
):
    add_bullet(doc, text)

add_heading(doc, "目前實現邊界", level=2)
for text in (
    "「忘記密碼」目前僅提示聯絡系統管理員，尚未提供自助重設流程。",
    "房款憑證的「查看」按鈕與租金收入的「查看詳情／查看全部」目前以訊息提示為主，尚未進入完整明細頁。",
    "付款憑證「保存草稿」僅保留在目前頁面狀態；重新整理後不會恢復。",
    "我的房產清單中的直接上傳入口目前導向文件資料模組；完整付款憑證流程由房款進度頁進入。",
):
    add_bullet(doc, text)

pages = [
    {
        "number": 1,
        "title": "登入頁",
        "intro": "作為 CCPS 房產管理工作區入口，兼顧品牌形象、帳號驗證與工作階段管理。",
        "display": [
            "左側呈現 CCPS 品牌、系統定位與 24/7、360°、100% 三項營運價值；右側為登入表單。",
            "表單包含帳號或電郵、密碼、顯示／隱藏密碼、記住我、忘記密碼及示範帳號提示。",
        ],
        "functions": [
            "登入前檢查必填欄位，提交期間鎖定按鈕並顯示處理狀態；失敗時於表單內呈現錯誤訊息。",
            "支援以使用者名稱或電郵登入；後端確認帳號為啟用狀態並以 BCrypt 驗證密碼。",
            "一般登入工作階段為 30 分鐘；啟用「記住我」後延長為 7 天，並於重新開啟頁面時檢查現有 Session。",
            "已實作登出功能：清除伺服器 Session 與瀏覽器登入資料，返回登入頁。",
        ],
        "note": "「忘記密碼」目前顯示聯絡系統管理員的提示，尚未串接電郵驗證或密碼重設。",
    },
    {
        "number": 2,
        "title": "我的房產",
        "intro": "集中呈現業主持有房產、付款概況、通知與待處理事項，作為登入後的主要總覽頁。",
        "display": [
            "上方 4 張摘要卡顯示名下房產數、本月租金收入、未繳房款與預備金餘額。",
            "房產清單顯示建案、單位、城市、總價、已繳／剩餘金額、期數、下一期日期與付款狀態；右側呈現最新通知及待處理事項。",
        ],
        "functions": [
            "可依建案名稱／單位編號搜尋，搭配建案及付款狀態篩選；提供查詢、重置、清單／精簡顯示切換。",
            "每頁顯示 3 筆房產並支援分頁；資料變動時自動校正頁碼。",
            "「查看詳情」開啟房產資訊視窗；通知與待處理卡可跳轉至通知中心或文件資料。",
            "資料由登入業主的儀表板 API 載入，並提供載入中、失敗提示及重新載入。",
        ],
        "note": "付款狀態統一對應正常繳費中、已繳清、即將到期、已逾期及待補繳款資料；直接上傳按鈕目前先導向文件資料模組。",
    },
    {
        "number": 3,
        "title": "房款進度詳情",
        "intro": "以單一房產為主軸，完整呈現購屋總價、分期計畫、到期狀態及付款憑證入口。",
        "display": [
            "房產卡顯示建案／單位、業主、電話、簽約日期及房款狀態；房產下拉選單可切換不同單位。",
            "頁面包含 7 項統計、付款里程碑、分期明細，以及右側已繳比例、最近付款與財務確認摘要。",
        ],
        "functions": [
            "切換房產後即時向後端載入付款計畫；支援載入中、失敗重試、無房產及未設定計畫等狀態。",
            "系統依應繳、已繳與日期計算分期狀態、下一期金額／日期，並顯示到期天數、已繳比例與剩餘期數。",
            "點擊「上傳付款憑證」時，優先帶入逾期、目前應繳或仍有未繳金額的分期，轉入上傳頁。",
        ],
        "note": "分期及最新付款資料已由 API 提供；「查看憑證」目前顯示功能接入提示，尚未開啟實際附件預覽。",
    },
    {
        "number": 4,
        "title": "上傳付款憑證",
        "intro": "讓業主針對指定房產與分期提交付款資訊及附件，進入財務審核流程。",
        "display": [
            "頂部摘要列顯示建案、單位、目前期數／階段、應繳金額、到期日及待上傳狀態。",
            "主表單包含付款金額／日期／方式、銀行、交易參考號、付款人、備註與附件；右側顯示三階段審核流程和提交記錄。",
        ],
        "functions": [
            "由房款進度頁自動帶入分期、金額、到期日及付款人；支援拖曳或點擊選檔、圖片預覽、PDF 標識及移除附件。",
            "前後端共同驗證必填欄位、付款金額、檔案格式／大小／數量；重複檔案不重複加入。",
            "提交後建立財務交易、文件記錄、文件關聯、付款收據及分期沖抵，並以待確認狀態送交財務。",
            "附件以安全檔名儲存並計算 SHA-256；成功後顯示收據編號、提交時間與頁面內提交記錄。",
        ],
        "note": "接受 JPG、PNG、PDF；1–4 個檔案、單檔不超過 10MB。付款日期不可晚於今天，金額不可高於可提交的剩餘款，備註上限 200 字。",
    },
    {
        "number": 5,
        "title": "租金收入",
        "intro": "提供業主查看指定月份的應收、已收、未收與年度累計，並追蹤租客付款及財務確認狀態。",
        "display": [
            "摘要卡呈現本月應收、本月已收、未收金額與年度累計，並顯示相較上月／去年之變動。",
            "租金明細包含房產、租客、應收日、應收／已收／未收、到帳日、狀態及財務確認；右側呈現年度趨勢與最近到帳。",
        ],
        "functions": [
            "可依房產、月份、年份及收款狀態向後端重新查詢；提供重置、載入、錯誤及無資料提示。",
            "表格支援分頁與每頁 8／10／20 筆切換，狀態涵蓋已收款、部分收款、未收款及已逾期。",
            "可將目前查詢結果匯出為 UTF-8 CSV，包含房產、單位、租客、金額、日期、狀態與財務確認欄位。",
            "年度趨勢依收入最高值換算長條比例，最近到帳列出房產、租客、日期與金額。",
        ],
        "note": "租金資料依登入業主查詢，篩選條件由 API 處理；頁面標示所有時間以馬來西亞 GMT+8 為準。",
    },
]

for page in pages:
    start_new_page_section(doc)
    compact_page = page["number"] in (3, 4)
    add_heading(doc, f"{page['number']:02d}　{page['title']}", level=1)
    add_paragraph(doc, page["intro"], size=11, color=MUTED, after=6)
    add_screenshot(
        doc,
        page["number"],
        f"{page['title']}頁面畫面",
        f"CCPS {page['title']}頁面截圖",
        width=5.55 if compact_page else 6.15,
    )
    add_heading(doc, "頁面展示內容", level=2)
    for item in page["display"]:
        add_bullet(doc, item, compact=compact_page)
    add_heading(doc, "已實現功能", level=2)
    for item in page["functions"]:
        add_bullet(doc, item, compact=compact_page)
    add_labeled_note(doc, "資料與實現說明", page["note"], compact=compact_page)

doc.core_properties.title = "CCPS 今日頁面功能交付說明"
doc.core_properties.subject = "頁面展示內容、已實現功能與後端整合"
doc.core_properties.author = "CCPS"
doc.core_properties.keywords = "CCPS, 房產管理, Vue 3, 功能交付, 頁面說明"
doc.core_properties.comments = "由原始頁面截圖及目前前後端程式整理。"

doc.save(OUTPUT)
print(OUTPUT)
